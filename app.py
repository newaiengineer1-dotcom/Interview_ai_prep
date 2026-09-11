import os
import io
import json
import re
import tempfile
from datetime import datetime

import streamlit as st
from groq import Groq
from pypdf import PdfReader
from docx import Document
from fpdf import FPDF

APP_NAME = "JOBLANDER AI"
DEFAULT_MODEL = "openai/gpt-oss-120b"
DEFAULT_STT = "whisper-large-v3-turbo"

st.set_page_config(
    page_title="JOBLANDER AI — Real-Time Interview Copilot",
    page_icon="🎙️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
:root { --bg:#07111f; --card:#0d1b2e; --line:#20344d; --text:#eef5ff; --muted:#9db0c8; --accent:#65d9ff; }
.stApp { background: radial-gradient(circle at 15% 0%, #122b49 0%, #07111f 42%, #050b14 100%); color:var(--text); }
.block-container { max-width: 1500px; padding-top: 1.3rem; }
.hero { padding:22px 24px; border:1px solid var(--line); border-radius:20px; background:linear-gradient(135deg,rgba(18,42,70,.9),rgba(8,18,32,.85)); box-shadow:0 12px 45px rgba(0,0,0,.22); }
.hero h1 { margin:0; font-size:34px; letter-spacing:.4px; }
.hero p { margin:.35rem 0 0; color:var(--muted); }
.card { padding:18px; border:1px solid var(--line); border-radius:18px; background:rgba(13,27,46,.82); margin-bottom:14px; }
.metric { font-size:30px; font-weight:800; }
.label { color:var(--muted); font-size:12px; text-transform:uppercase; letter-spacing:1px; }
.question { font-size:24px; font-weight:700; line-height:1.35; }
.answer { font-size:17px; line-height:1.6; }
.badge { display:inline-block; padding:5px 10px; border-radius:999px; background:#132941; border:1px solid #2a4664; margin:3px; font-size:12px; }
.good { color:#77e3a5; } .warn { color:#ffd37a; } .bad { color:#ff8d9c; }
.small { color:var(--muted); font-size:13px; }
div[data-testid="stFileUploader"] { border-radius:14px; }
button[kind="primary"] { border-radius:12px; }
</style>
""", unsafe_allow_html=True)

def extract_pdf(file):
    reader = PdfReader(file)
    return "\n".join((p.extract_text() or "") for p in reader.pages)

def extract_docx(file):
    doc = Document(file)
    return "\n".join(p.text for p in doc.paragraphs)

def extract_text(file):
    name = file.name.lower()
    if name.endswith(".pdf"):
        return extract_pdf(file)
    if name.endswith(".docx"):
        return extract_docx(file)
    return file.getvalue().decode("utf-8", errors="ignore")

def groq_client(key):
    return Groq(api_key=key.strip())

def chat_json(client, model, system, user):
    response = client.chat.completions.create(
        model=model,
        temperature=0.15,
        response_format={"type": "json_object"},
        messages=[{"role":"system","content":system},{"role":"user","content":user}],
    )
    return json.loads(response.choices[0].message.content)

def grounded_system():
    return """You are JOBLANDER AI, a professional interview copilot.
Use ONLY the supplied CV, JD, user facts, and explicitly supplied research.
Never invent candidate employers, projects, responsibilities, numbers,
certifications, achievements, technologies, or experience.
If evidence is absent, mark it UNKNOWN. If a candidate claim is unsupported,
flag it as UNSUPPORTED rather than accepting it.
Questions must be concise, practical, human, role-specific and conversational.
Answers must be short, natural, first-person and speakable.
Do not write generic textbook questions when CV/JD evidence is available.
Return valid JSON only."""

def generate_profile(client, model, cv, jd, role):
    system = grounded_system()
    user = f"""Create a verified interview profile.
TARGET ROLE:
{role}

CV:
{cv[:30000]}

JOB DESCRIPTION:
{jd[:30000]}

Return JSON with:
summary, verified_skills, verified_experience, jd_requirements,
match_score_0_100, gaps, risk_flags."""
    return chat_json(client, model, system, user)

def generate_question(client, model, cv, jd, profile, role, mode, difficulty, history):
    system = grounded_system()
    user = f"""Generate ONE next interview question.
Role: {role}
Mode: {mode}
Difficulty: {difficulty}
Verified profile:
{json.dumps(profile, ensure_ascii=False)}

CV:
{cv[:18000]}

JD:
{jd[:18000]}

Previous interview:
{json.dumps(history[-8:], ensure_ascii=False)}

Rules:
- 1 short question, normally under 25 words.
- Prefer a realistic follow-up based on the previous answer.
- If verifying CV evidence, ask the candidate to explain their actual role.
- Do not assume missing facts.
Return:
{{"question":"...", "type":"technical|behavioral|resume|scenario|follow_up",
"evidence_basis":"...", "why_this_question":"..."}}"""
    return chat_json(client, model, system, user)

def analyze_answer(client, model, cv, jd, profile, question, answer):
    system = grounded_system()
    user = f"""Evaluate the candidate's answer.

QUESTION:
{question}

ANSWER:
{answer}

VERIFIED PROFILE:
{json.dumps(profile, ensure_ascii=False)}

CV:
{cv[:18000]}

JD:
{jd[:18000]}

Return JSON:
score_0_100, relevance_0_100, completeness_0_100,
communication_0_100, evidence_grounding_0_100,
strengths (array), missing_points (array), unsupported_claims (array),
short_improved_answer, verdict, follow_up_focus.
The improved answer must NOT add facts absent from evidence."""
    return chat_json(client, model, system, user)

def transcribe(client, audio_bytes, filename, model):
    suffix = os.path.splitext(filename)[1] or ".wav"
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
        f.write(audio_bytes)
        path = f.name
    try:
        with open(path, "rb") as audio_file:
            result = client.audio.transcriptions.create(
                file=(os.path.basename(path), audio_file.read()),
                model=model,
                response_format="json",
            )
        return result.text
    finally:
        try: os.unlink(path)
        except OSError: pass

def make_pdf(session, role, profile):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(0, 12, "JOBLANDER AI - Interview Report", ln=True)
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 7, datetime.now().strftime("%Y-%m-%d %H:%M"), ln=True)
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13); pdf.cell(0, 8, "Target Role", ln=True)
    pdf.set_font("Helvetica", "", 10); pdf.multi_cell(0, 6, role)
    pdf.set_font("Helvetica", "B", 13); pdf.cell(0, 8, "JD/CV Match", ln=True)
    pdf.set_font("Helvetica", "", 10); pdf.multi_cell(0, 6, str(profile.get("match_score_0_100","N/A")))
    for i, item in enumerate(session):
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 6, f"Q{i+1}. {item['question']}")
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, f"Candidate: {item['answer']}")
        pdf.multi_cell(0, 6, f"Score: {item['analysis'].get('score_0_100','N/A')}")
        pdf.multi_cell(0, 6, "Feedback: " + "; ".join(item['analysis'].get('missing_points', [])))
        improved = item['analysis'].get('short_improved_answer','')
        if improved:
            pdf.multi_cell(0, 6, "Improved: " + improved)
    return bytes(pdf.output())

if "profile" not in st.session_state: st.session_state.profile = None
if "history" not in st.session_state: st.session_state.history = []
if "question" not in st.session_state: st.session_state.question = None
if "last_analysis" not in st.session_state: st.session_state.last_analysis = None

st.markdown("""<div class="hero">
<h1>🎙️ JOBLANDER AI</h1>
<p>Real-Time, Evidence-Grounded Interview Copilot — CV + JD + Voice + Text + Adaptive Follow-Ups</p>
</div>""", unsafe_allow_html=True)
st.write("")

with st.sidebar:
    st.markdown("### ⚙️ Interview Setup")
    api_key = st.text_input("Groq API Key", type="password", help="Use Streamlit secrets in deployment when preferred.")
    model = st.text_input("LLM Model", value=DEFAULT_MODEL)
    stt_model = st.selectbox("Speech-to-Text", [DEFAULT_STT, "whisper-large-v3"])
    role = st.text_input("Target Role", placeholder="e.g. Lender Technical Advisor - Solar + BESS")
    company = st.text_input("Company (optional)")
    mode = st.selectbox("Interview Mode", ["Full Interview","Technical","Behavioral / HR","Hiring Manager","Resume Verification","Stress / Follow-up"])
    difficulty = st.selectbox("Difficulty", ["Intermediate","Advanced","Expert","Executive"])
    web_research = st.toggle("Web research", value=False, help="This package keeps research opt-in. Add a search provider/API before enabling automated external research.")
    st.caption("Evidence rule: the copilot must not invent candidate facts.")
    if st.button("Reset Session", use_container_width=True):
        st.session_state.profile = None
        st.session_state.history = []
        st.session_state.question = None
        st.session_state.last_analysis = None
        st.rerun()

c1, c2 = st.columns(2)
with c1:
    st.markdown('<div class="card"><div class="label">CV / Resume</div>', unsafe_allow_html=True)
    cv_file = st.file_uploader("Upload CV / Resume", type=["pdf","docx","txt"], key="cv")
    st.markdown("</div>", unsafe_allow_html=True)
with c2:
    st.markdown('<div class="card"><div class="label">Job Description</div>', unsafe_allow_html=True)
    jd_file = st.file_uploader("Upload Job Description", type=["pdf","docx","txt"], key="jd")
    st.markdown("</div>", unsafe_allow_html=True)

if cv_file and jd_file and api_key and role:
    if st.button("🚀 Build Evidence-Grounded Interview", type="primary", use_container_width=True):
        try:
            with st.spinner("Reading CV and JD and building verified interview context..."):
                cv_text, jd_text = extract_text(cv_file), extract_text(jd_file)
                client = groq_client(api_key)
                st.session_state.profile = generate_profile(client, model, cv_text, jd_text, role)
                st.session_state.cv_text, st.session_state.jd_text = cv_text, jd_text
                st.session_state.history = []
                st.session_state.last_analysis = None
                st.session_state.question = generate_question(
                    client, model, cv_text, jd_text, st.session_state.profile,
                    role, mode, difficulty, []
                )
        except Exception as e:
            st.error(f"Could not initialize interview: {e}")

if st.session_state.profile:
    profile = st.session_state.profile
    st.write("")
    a,b,c,d = st.columns(4)
    with a:
        st.markdown('<div class="card"><div class="label">JD / CV Match</div><div class="metric">{}%</div></div>'.format(profile.get("match_score_0_100","—")), unsafe_allow_html=True)
    with b:
        st.markdown('<div class="card"><div class="label">Questions</div><div class="metric">{}</div></div>'.format(len(st.session_state.history)), unsafe_allow_html=True)
    with c:
        last = st.session_state.last_analysis or {}
        st.markdown('<div class="card"><div class="label">Last Answer</div><div class="metric">{}%</div></div>'.format(last.get("score_0_100","—")), unsafe_allow_html=True)
    with d:
        st.markdown('<div class="card"><div class="label">Mode</div><div class="metric" style="font-size:18px">{}</div></div>'.format(mode), unsafe_allow_html=True)

    left, center, right = st.columns([1.0, 2.0, 1.0])

    with left:
        st.markdown('<div class="card"><b>Verified Skills</b><br>'.format(), unsafe_allow_html=True)
        for x in profile.get("verified_skills", [])[:12]:
            st.markdown(f'<span class="badge">{x}</span>', unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)
        st.markdown('<div class="card"><b>Gaps / Unknowns</b>', unsafe_allow_html=True)
        for x in profile.get("gaps", [])[:8]:
            st.markdown(f"<p class='small'>• {x}</p>", unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)

    with center:
        st.markdown('<div class="card"><div class="label">Interviewer Question</div>', unsafe_allow_html=True)
        q = st.session_state.question or {}
        st.markdown(f"<div class='question'>{q.get('question','Ready')}</div>", unsafe_allow_html=True)
        st.caption(f"Type: {q.get('type','')}  •  Evidence: {q.get('evidence_basis','')}")
        st.markdown("</div>", unsafe_allow_html=True)

        tab_voice, tab_text = st.tabs(["🎙 Voice Answer", "⌨ Type Answer"])
        answer = ""
        with tab_voice:
            audio = st.audio_input("Record your answer")
            if audio:
                if st.button("Transcribe & Analyze Voice", type="primary"):
                    try:
                        with st.spinner("Transcribing your answer..."):
                            answer = transcribe(groq_client(api_key), audio.getvalue(), "answer.wav", stt_model)
                            st.session_state.pending_answer = answer
                    except Exception as e:
                        st.error(f"Voice transcription failed: {e}")
            if st.session_state.get("pending_answer"):
                st.text_area("Detected speech", value=st.session_state.pending_answer, height=120, key="voice_text")
                if st.button("Analyze Detected Answer", key="analyze_voice"):
                    answer = st.session_state.pending_answer
                    st.session_state.analyze_now = answer
        with tab_text:
            typed = st.text_area("Your answer", height=150, placeholder="Type a short, natural interview answer...")
            if st.button("Analyze Typed Answer", type="primary"):
                st.session_state.analyze_now = typed

        if st.session_state.get("analyze_now"):
            try:
                answer = st.session_state.pop("analyze_now")
                if not answer.strip():
                    st.warning("Please provide an answer.")
                else:
                    with st.spinner("Checking relevance, completeness and evidence grounding..."):
                        analysis = analyze_answer(
                            groq_client(api_key), model, st.session_state.cv_text,
                            st.session_state.jd_text, profile, q.get("question",""), answer
                        )
                    st.session_state.last_analysis = analysis
                    st.session_state.history.append({
                        "question": q.get("question",""),
                        "answer": answer,
                        "analysis": analysis,
                    })
                    st.session_state.question = generate_question(
                        groq_client(api_key), model, st.session_state.cv_text,
                        st.session_state.jd_text, profile, role, mode, difficulty,
                        st.session_state.history
                    )
                    st.rerun()
            except Exception as e:
                st.error(f"Analysis failed: {e}")

    with right:
        st.markdown('<div class="card"><div class="label">AI Copilot Feedback</div>', unsafe_allow_html=True)
        an = st.session_state.last_analysis
        if an:
            score = an.get("score_0_100","—")
            st.markdown(f"<div class='metric'>{score}%</div>", unsafe_allow_html=True)
            st.markdown(f"**Verdict:** {an.get('verdict','')}")
            st.markdown("**Strengths**")
            for x in an.get("strengths", []): st.markdown(f"• {x}")
            st.markdown("**Add / Improve**")
            for x in an.get("missing_points", []): st.markdown(f"• {x}")
            if an.get("unsupported_claims"):
                st.markdown("**⚠ Unsupported claims**")
                for x in an["unsupported_claims"]: st.markdown(f"• {x}")
            st.markdown("**Short interview-ready answer**")
            st.info(an.get("short_improved_answer",""))
        else:
            st.markdown("<p class='small'>Answer a question to see grounded feedback.</p>", unsafe_allow_html=True)
        st.markdown("</div>", unsafe_allow_html=True)

    st.divider()
    st.subheader("📄 Interview Report")
    if st.session_state.history:
        report_pdf = make_pdf(st.session_state.history, role, profile)
        st.download_button("⬇️ Download Final Interview Report (PDF)", report_pdf, "JobLander_Interview_Report.pdf", "application/pdf")
        st.markdown("### Session History")
        for i, item in enumerate(reversed(st.session_state.history), 1):
            with st.expander(f"Q{len(st.session_state.history)-i+1}: {item['question']}"):
                st.write("**Candidate:**", item["answer"])
                st.write("**Score:**", item["analysis"].get("score_0_100"))
                st.write("**Feedback:**", item["analysis"].get("missing_points", []))
else:
    st.info("Enter the Groq key, target role, CV and Job Description, then click **Build Evidence-Grounded Interview**.")
